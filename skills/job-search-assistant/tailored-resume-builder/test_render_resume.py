"""Tests for render_resume.py and the style-file parser.
Run with: python3 -m pytest -q skills/job-search-assistant/tailored-resume-builder

Uses default-style.md and the synthetic content in ../evals/fixtures; variants of the style are
written to a temp folder. The over-length test needs LibreOffice and skips without it.
"""

import json
import subprocess
import sys
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.oxml.ns import qn  # noqa: E402

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from render_resume import StyleError, bbox_fill, page_warnings, parse_style  # noqa: E402

DEFAULT_STYLE = (HERE / "default-style.md").read_text(encoding="utf-8")
CONTENT = json.loads((HERE.parent / "evals" / "fixtures" / "synthetic_content.json").read_text(encoding="utf-8"))


def render(tmp_path, style=DEFAULT_STYLE, content=None, args=()):
    s, c, out = tmp_path / "style.md", tmp_path / "content.json", tmp_path / "out.docx"
    s.write_text(style, encoding="utf-8")
    c.write_text(json.dumps(content or CONTENT), encoding="utf-8")
    r = subprocess.run([sys.executable, str(HERE / "render_resume.py"), "--style", str(s),
                        "--content", str(c), "--out", str(out), *args],
                       capture_output=True, text=True, encoding="utf-8")
    return r, (docx.Document(out) if r.returncode == 0 else None)


def edited(fn):
    c = json.loads(json.dumps(CONTENT))
    fn(c)
    return c


def with_line(key, value, style=DEFAULT_STYLE):
    lines = [f"- {key}: {value}" if line.startswith(f"- {key}:") else line for line in style.splitlines()]
    return "\n".join(lines)


def texts(d):
    return [p.text for p in d.paragraphs]


def para(d, text):
    return next(p for p in d.paragraphs if p.text.startswith(text))


# ---------------------------------------------------------------- the default style


def test_default_style_parses():
    st = parse_style(DEFAULT_STYLE)
    assert st.size == "letter" and st.competencies == ("grid", 3) and st.dates == "right"
    assert st.t("Role").font == "Calibri"  # inherited from Body


def test_renders_page_and_header(tmp_path):
    r, d = render(tmp_path)
    assert r.returncode == 0, r.stdout
    sec = d.sections[0]
    assert round(sec.page_width.inches, 2) == 8.5 and round(sec.left_margin.inches, 2) == 0.6
    name = d.paragraphs[0].runs[0]
    assert name.text == "JANE EXAMPLE" and name.font.size.pt == 20 and name.font.bold and name.font.all_caps
    assert str(name.font.color.rgb) == "1F3864"
    assert d.paragraphs[2].text == " · ".join(CONTENT["contact"])


def test_sections_follow_the_style_order(tmp_path):
    _, d = render(tmp_path)
    t = texts(d)
    order = [t.index(x) for x in ("Career Highlights", "Professional Experience", "Education", "Certifications")]
    assert order == sorted(order)
    assert "Summary" not in t  # summary title is empty in the default style


def test_competencies_grid_and_bullets(tmp_path):
    _, d = render(tmp_path)
    (table,) = d.tables
    assert len(table.columns) == 3
    cell_items = [p.text for cell in table.rows[0].cells for p in cell.paragraphs if p.text]
    assert cell_items == CONTENT["competencies"]
    highlight = para(d, CONTENT["highlights"][0])
    assert highlight._p.pPr.numPr is not None


def test_dates_right_aligned_and_heading_rule(tmp_path):
    _, d = render(tmp_path)
    company = para(d, "Example Corp")
    assert "\t2020 - Present" in company.text
    assert company.paragraph_format.tab_stops[0].position.inches == pytest.approx(8.5 - 1.2, abs=0.01)
    heading = para(d, "Professional Experience")
    assert heading._p.pPr.find(qn("w:pBdr")) is not None


def test_bullet_character_comes_from_the_style(tmp_path):
    _, d = render(tmp_path, with_line("Bullet", "▪"))
    numbering = d.part.numbering_part.element
    assert any(el.get(qn("w:val")) == "▪" for el in numbering.iter(qn("w:lvlText")))


# ---------------------------------------------------------------- variants


def test_inline_competencies_inline_dates_a4_centered(tmp_path):
    style = with_line("Competencies", 'inline " | "')
    style = with_line("Dates", "inline", style)
    style = with_line("Size", "A4", style)
    style = with_line("Header alignment", "center", style)
    _, d = render(tmp_path, style)
    assert not d.tables
    assert " | ".join(CONTENT["competencies"]) in texts(d)
    assert para(d, "Example Corp").text.endswith(" | 2020 - Present")
    assert round(d.sections[0].page_width.inches, 2) == 8.27
    assert d.paragraphs[0].alignment == 1  # center


def test_section_order_and_titles_are_honoured(tmp_path):
    style = with_line("Section order", "experience, education")
    style = with_line("Section titles", "experience = Work History; education = Schooling", style)
    _, d = render(tmp_path, style)
    t = texts(d)
    assert "Work History" in t and "Schooling" in t
    assert "Career Highlights" not in t and not d.tables
    assert t.index("Work History") < t.index("Schooling")


def test_accomplishments_label(tmp_path):
    _, d = render(tmp_path, with_line("Accomplishments label", "Key Accomplishments:"))
    assert texts(d).count("Key Accomplishments:") == 2  # one per role with bullets


@pytest.mark.parametrize("key, value, message", [
    ("Size", "Legal", "Letter or A4"),
    ("Body", "Calibri", "font and a size"),
    ("Name", "Calibri, 20, color navy", "6 hex digits"),
    ("Role", "Calibri, rule", "only applies to Section headings"),
    ("Competencies", "grid 7", "grid N"),
    ("Section order", "summary, hobbies", "sections must be distinct"),
    ("Margins (in)", "top 0.5, left 0.5", "all four"),
])
def test_bad_style_values_fail_with_the_line(tmp_path, key, value, message):
    r, _ = render(tmp_path, with_line(key, value))
    assert r.returncode == 2
    assert message in json.loads(r.stdout)["error"]


def test_unknown_key_and_missing_version_fail(tmp_path):
    r, _ = render(tmp_path, DEFAULT_STYLE.replace("- Bullet: •", "- Bullets please: •"))
    assert r.returncode == 2 and "unknown Layout setting" in r.stdout
    r, _ = render(tmp_path, DEFAULT_STYLE.replace("Style format: 1", "Style format: 9"))
    assert r.returncode == 2 and "Style format: 1" in r.stdout
    with pytest.raises(StyleError):
        parse_style(DEFAULT_STYLE.replace("## Layout", "## Colours"))


def test_bad_content_fails(tmp_path):
    c = dict(CONTENT)
    c["competencies"] = [f"s{i}" for i in range(13)]
    r, _ = render(tmp_path, content=c)
    assert r.returncode == 2 and "too many competencies" in r.stdout


# ---------------------------------------------------------------- work history


def test_role_dates_and_place_shown_only_where_they_differ(tmp_path):
    _, d = render(tmp_path)
    director = para(d, "Director of IT")  # same place as the company, different dates
    assert director.text == "Director of IT\t2022 - Present"
    assert para(d, "Program Manager").text == "Program Manager\t2020 - 2022"


def test_single_role_dates_not_repeated(tmp_path):
    def single(c):
        c["experience"].append({"company": "Sample Systems", "location": "Dayton, OH", "dates": "2012 \u2013 2018",
                                "roles": [{"title": "IT Project Manager", "location": "Dayton, OH",
                                           "dates": "2012 - 2018", "bullets": ["Managed synthetic projects."]}]})
    _, d = render(tmp_path, content=edited(single))
    assert para(d, "Sample Systems").text == "Sample Systems, Dayton, OH\t2012 \u2013 2018"
    assert para(d, "IT Project Manager").text == "IT Project Manager"
    assert sum("2012" in t for t in texts(d)) == 1


def test_company_description_rejected(tmp_path):
    r, _ = render(tmp_path, content=edited(lambda c: c["experience"][0].update(description="A fictional company.")))
    assert r.returncode == 2 and "company descriptions" in json.loads(r.stdout)["error"]


def test_every_role_needs_a_line(tmp_path):
    r, _ = render(tmp_path, content=edited(lambda c: c["experience"][0]["roles"][1].update(bullets=[])))
    assert r.returncode == 2 and "needs at least one line" in json.loads(r.stdout)["error"]

    def summary_only(c):  # a summary with no bullets key at all is fine
        c["experience"][0]["roles"][1].pop("bullets")
        c["experience"][0]["roles"][1]["summary"] = "Delivered synthetic projects."
    _, d = render(tmp_path, with_line("Accomplishments label", "Key Accomplishments:"), edited(summary_only))
    assert "Delivered synthetic projects." in texts(d)
    assert texts(d).count("Key Accomplishments:") == 1  # only the role with bullets gets the label


# ---------------------------------------------------------------- page limit and last-page fill

BBOX = """<html><body><doc>
  <page width="612.000000" height="792.000000">
    <word xMin="43.2" yMin="36.0" xMax="90.0" yMax="48.0">First</word>
    <word xMin="43.2" yMin="200.0" xMax="90.0" yMax="{bottom}">last</word>
  </page>
</doc></body></html>"""
MARGINS = dict(top=0.5, bottom=0.5, left=0.6, right=0.6)  # 36pt top and bottom: 720pt of text area


@pytest.mark.parametrize("bottom, fill", [(36 + 72, 0.1), (36 + 360, 0.5), (36 + 720, 1.0), (790, 1.0)])
def test_bbox_fill(bottom, fill):
    assert bbox_fill(BBOX.format(bottom=bottom), MARGINS) == fill


def test_bbox_fill_without_words_or_page():
    assert bbox_fill("<page width=\"612\" height=\"792\"></page>", MARGINS) is None
    assert bbox_fill("", MARGINS) is None


@pytest.mark.parametrize("pages, max_pages, fill, expected", [
    (2, 3, None, []),
    (2, 3, 0.1, []),  # a short 2nd page is fine
    (3, 3, 0.5, []),
    (3, 3, 0.8, []),
    (3, 3, 0.49, ["page 3 is only 49% full (needs 50%+); trim to 2 pages"]),
    (3, 3, None, ["page 3 fill couldn't be measured; trim to 2 pages"]),
    (3, 2, 0.9, ["resume is 3 pages; the limit is 2"]),
    (4, 3, 0.9, ["resume is 4 pages; the limit is 3"]),
    (2, 1, 0.9, ["resume is 2 pages; the limit is 1"]),
    (None, 3, None, ["could not render to PDF; page count unverified"]),
])
def test_page_warnings(pages, max_pages, fill, expected):
    assert page_warnings(pages, max_pages, fill) == expected


@pytest.mark.parametrize("value", ["0", "4"])
def test_max_pages_out_of_range(tmp_path, value):
    r, _ = render(tmp_path, args=("--max-pages", value))
    assert r.returncode == 2 and "max pages must be 1, 2 or 3" in json.loads(r.stdout)["error"]


def long_content(bullets: int) -> dict:
    return edited(lambda c: c["experience"][0]["roles"][0].update(
        bullets=[f"Ran synthetic program number {i} to plan." for i in range(bullets)]))


def test_page_count_rules(tmp_path):
    r, _ = render(tmp_path)
    res = json.loads(r.stdout)
    if res["pages"] is None:
        assert res["warnings"] == ["could not render to PDF; page count unverified"]
        pytest.skip("LibreOffice not installed; page count can't be checked")
    assert res["pages"] <= 2 and res["warnings"] == [] and res["max_pages"] == 3
    r, _ = render(tmp_path, content=long_content(400))
    res = json.loads(r.stdout)
    assert res["pages"] > 3 and res["warnings"] == [f"resume is {res['pages']} pages; the limit is 3"]
    r, _ = render(tmp_path, content=long_content(400), args=("--max-pages", "2"))
    assert any("the limit is 2" in w for w in json.loads(r.stdout)["warnings"])


def test_third_page_fill(tmp_path):
    """Grow the resume until it just spills onto page 3 (under half full: trim warning), then
    until page 3 is past half (no warning). Steps, not exact counts, so fonts don't matter."""
    def build(n):
        r, _ = render(tmp_path, content=long_content(n))
        return json.loads(r.stdout)
    n, res = 60, build(60)
    if res["pages"] is None or res["last_page_fill"] is None:
        pytest.skip("LibreOffice or pdftotext not installed; page fill can't be checked")
    while res["pages"] < 3 and n < 400:
        n += 4
        res = build(n)
    assert res["pages"] == 3 and res["last_page_fill"] < 0.5
    assert any("trim to 2 pages" in w for w in res["warnings"])
    while res["pages"] == 3 and res["last_page_fill"] < 0.5 and n < 400:
        n += 4
        res = build(n)
    assert res["pages"] == 3 and res["last_page_fill"] >= 0.5 and res["warnings"] == []
