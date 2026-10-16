"""Tests for extract_style.py, which drafts a resume style file from a Word template.
Run with: python3 -m pytest -q skills/job-search-assistant/tailored-resume-builder

Templates are generated here: one rendered by render_resume.py from default-style.md (a round
trip), one built with direct formatting only, and one with features the renderer can't carry over.
"""

import base64
import json
import subprocess
import sys
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from extract_style import extract  # noqa: E402
from render_resume import parse_style  # noqa: E402

CONTENT = json.loads((HERE.parent / "evals" / "fixtures" / "synthetic_content.json").read_text(encoding="utf-8"))
CONTENT["summary"] = ("IT program delivery leader with 14 years in IT, the last 9 leading cloud and "
                      "infrastructure programs for a logistics company.")
PNG_1PX = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def rendered_template(tmp_path, style_text=None):
    style = tmp_path / "style.md"
    style.write_text(style_text or (HERE / "default-style.md").read_text(encoding="utf-8"), encoding="utf-8")
    content = tmp_path / "content.json"
    content.write_text(json.dumps(CONTENT), encoding="utf-8")
    out = tmp_path / "template.docx"
    subprocess.run([sys.executable, str(HERE / "render_resume.py"), "--style", str(style), "--content", str(content),
                    "--out", str(out)], check=True, capture_output=True)
    return out


def direct_formatting_template(tmp_path, extras=False):
    """A template with no named styles: everything is set on the runs, dates are inline."""
    d = docx.Document()
    def line(text, size=11, bold=False, italic=False, color=None, center=False):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold, r.italic = "Georgia", Pt(size), bold or None, italic or None
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if center:
            p.alignment = 1
        return p
    line("SAM SAMPLE", 24, bold=True, color="7A1F1F", center=True)
    line("OPERATIONS LEADER", 12, bold=True, center=True)
    line("555-0000 | Sample City | sam@example.com", 9, center=True)
    line("Operations leader with twelve years of experience running plants, budgets and teams across two regions.")
    line("EXPERIENCE", 12, bold=True, color="7A1F1F")
    line("Sample Industries | 2019 - Present", bold=True)
    line("A fictional manufacturer.", italic=True)
    line("Plant Manager | 2019 - Present", bold=True)
    line("Ran a fictional plant of 200 people with a large budget and many projects.")
    line("SKILLS", 12, bold=True, color="7A1F1F")
    line("Budgeting | Lean | Safety | Hiring")
    line("EDUCATION", 12, bold=True, color="7A1F1F")
    line("B.S., Sample University")
    if extras:
        img = tmp_path / "dot.png"
        img.write_bytes(PNG_1PX)
        d.add_picture(str(img), width=Inches(0.2))
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        d.sections[0]._sectPr.append(cols)
        d.sections[0].header.paragraphs[0].text = "Confidential"
    out = tmp_path / ("extras.docx" if extras else "direct.docx")
    d.save(out)
    return out


def lines(style_text):
    return {line[2:].split(":", 1)[0]: line.split(":", 1)[1].strip()
            for line in style_text.splitlines() if line.startswith("- ") and ":" in line}


# ---------------------------------------------------------------- round trip


def test_round_trip_recovers_the_default_style(tmp_path):
    style, report = extract(str(rendered_template(tmp_path)))
    got = lines(style)
    assert got["Size"] == "Letter"
    assert got["Margins (in)"] == "top 0.5, bottom 0.5, left 0.6, right 0.6"
    assert got["Body"] == "Calibri, 10.5"
    assert got["Name"] == "20, bold, caps, color 1F3864"
    assert got["Section headings"] == "11, bold, caps, color 1F3864, rule"
    assert got["Company description"].startswith("italic")
    assert got["Location"] == "italic"
    assert got["Header order"] == "name, headline, contact"
    assert got["Section order"] == "summary, competencies, highlights, experience, education, certifications, additional"
    assert got["Section titles"].startswith("summary = ; competencies = Core Competencies; highlights = Career Highlights")
    assert got["Competencies"] == "grid 3"
    assert got["Bullet"] == "•"
    assert got["Dates"] == "right"
    assert report["not_carried_over"] == []


def test_round_trip_of_a_centred_layout(tmp_path):
    default = (HERE / "default-style.md").read_text(encoding="utf-8")
    variant = (default.replace("- Header alignment: left", "- Header alignment: center")
               .replace("- Header order: name, headline, contact", "- Header order: name, contact, headline")
               .replace("- Section heading alignment: left", "- Section heading alignment: center")
               .replace("- Accomplishments label:", "- Accomplishments label: Key Accomplishments:"))
    got = lines(extract(str(rendered_template(tmp_path, variant)))[0])
    assert got["Header alignment"] == "center"
    assert got["Header order"] == "name, contact, headline"
    assert got["Section heading alignment"] == "center"
    assert got["Accomplishments label"] == "Key Accomplishments:"


# ---------------------------------------------------------------- other templates


def test_direct_formatting_template(tmp_path):
    style, report = extract(str(direct_formatting_template(tmp_path)))
    got = lines(style)
    assert got["Body"] == "Georgia, 11"
    assert got["Name"] == "24, bold, color 7A1F1F"
    assert got["Header alignment"] == "center"
    assert got["Contact separator"] == '" | "'
    assert got["Dates"] == "inline"
    assert got["Section order"] == "summary, experience, competencies, education"
    assert got["Competencies"].startswith("inline")


def test_features_not_carried_over_are_reported(tmp_path):
    _, report = extract(str(direct_formatting_template(tmp_path, extras=True)))
    nco = " ".join(report["not_carried_over"])
    assert "images" in nco and "multi-column" in nco and "header or footer" in nco


def test_every_extracted_style_is_valid_for_the_renderer(tmp_path):
    templates = [rendered_template(tmp_path), direct_formatting_template(tmp_path),
                 direct_formatting_template(tmp_path, extras=True),
                 HERE.parent / "evals" / "fixtures" / "Old-Resume.docx"]
    for t in templates:
        parse_style(extract(str(t))[0])  # raises StyleError if invalid


def test_not_a_word_file_fails_cleanly(tmp_path):
    bad = tmp_path / "resume.docx"
    bad.write_text("not a zip", encoding="utf-8")
    r = subprocess.run([sys.executable, str(HERE / "extract_style.py"), str(bad)], capture_output=True, text=True)
    assert r.returncode == 2 and "couldn't read" in r.stdout


def test_cli_writes_the_style_and_reports(tmp_path):
    out = tmp_path / "resume-style.md"
    r = subprocess.run([sys.executable, str(HERE / "extract_style.py"), str(rendered_template(tmp_path)), "--out", str(out)],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0
    report = json.loads(r.stdout)
    assert report["style"] == str(out) and out.read_text(encoding="utf-8").startswith("# Resume style")
    assert {g["key"] for g in report["guesses"]} >= {"Name", "Body", "Competencies", "Dates", "Bullet"}
