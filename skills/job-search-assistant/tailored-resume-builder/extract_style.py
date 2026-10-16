#!/usr/bin/env python3
"""Draft a resume style file from a Word resume template (.docx).

Usage: extract_style.py template.docx [--out resume-style.md]

Reads the template's page setup and the effective formatting of its parts (name, headline,
contact line, section headings, company and role lines, dates, bullets, the competencies table)
and writes a style file in the format of resume-style-format.md, for render_resume.py.
Prints JSON: {"style": <path or null>, "guesses": [...], "not_carried_over": [...]}. Each guess
says what was found and how, with a confidence, so a person can confirm the uncertain ones.
Exit 2 with {"error": ...} if the file can't be read as a .docx.

It's a one-time conversion: the user checks a preview before the style file is saved.
"""

from __future__ import annotations

import argparse
import collections
import datetime
import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

YEAR_RANGE = re.compile(r"\b(19|20)\d{2}\b\s*[-–—to]+\s*(\b(19|20)\d{2}\b|present|current|now)", re.I)
SECTION_WORDS = [  # (section, pattern), first match wins
    ("competencies", r"competenc|skills|expertise|strengths"),
    ("highlights", r"highlight|achievement|accomplishments\b(?!:)|key results"),
    ("experience", r"experience|employment|work history|career history"),
    ("education", r"education"),
    ("certifications", r"certif|licens|training"),
    ("additional", r"additional|languages|interests|other information|volunteer"),
    ("summary", r"summary|profile|objective|about me"),
]
COMMON_FONTS = {
    "arial", "calibri", "calibri light", "cambria", "georgia", "garamond", "helvetica", "times new roman",
    "verdana", "tahoma", "trebuchet ms", "book antiqua", "palatino linotype", "century gothic", "segoe ui",
    "aptos", "carlito", "caladea", "liberation sans", "liberation serif", "dejavu sans", "dejavu serif",
}
# Bullet glyphs Word stores in symbol fonts' private-use area -> their ordinary characters
BULLET_MAP = {"": "•", "": "▪", "": "➢", "": "❖", "": "✓", "o": "◦", "§": "▪"}


class Template:
    def __init__(self, path: str):
        self.d = docx.Document(path)
        self.theme_fonts = self._theme_fonts()
        self.defaults = self._doc_defaults()

    # ------------------------------------------------------------ effective formatting
    def _theme_fonts(self) -> dict:
        fonts = {"minor": None, "major": None}
        for rel in self.d.part.rels.values():
            if rel.reltype.endswith("/theme"):
                xml = rel.target_part.blob.decode("utf-8", "ignore")
                for kind in fonts:
                    m = re.search(rf"<a:{kind}Font>\s*<a:latin typeface=\"([^\"]+)\"", xml)
                    fonts[kind] = m.group(1) if m else None
        return fonts

    def _font_from_rfonts(self, rfonts) -> str | None:
        if rfonts is None:
            return None
        name = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:hAnsi"))
        if name:
            return name
        theme = rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:hAnsiTheme"))
        if theme:
            return self.theme_fonts["major" if theme.startswith("major") else "minor"]
        return None

    def _doc_defaults(self) -> dict:
        out = {"font": None, "size": None}
        rpr = self.d.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            out["font"] = self._font_from_rfonts(rpr.find(qn("w:rFonts")))
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                out["size"] = int(sz.get(qn("w:val"))) / 2
        out["font"] = out["font"] or self.theme_fonts["minor"] or "Calibri"
        out["size"] = out["size"] or 11.0
        return out

    @staticmethod
    def _chain(style):
        while style is not None:
            yield style
            style = style.base_style

    def fmt(self, run, para) -> dict:
        """Effective font, size, bold, italic, caps, underline and colour of a run."""
        sources = [run.font]
        if run.style is not None and run.style.type == 2 and run.style.name != "Default Paragraph Font":
            sources += [s.font for s in self._chain(run.style)]
        sources += [s.font for s in self._chain(para.style)]
        out = {}
        for prop in ("bold", "italic", "all_caps", "underline"):
            out[prop] = next((getattr(f, prop) for f in sources if getattr(f, prop) is not None), False)
        out["size"] = next((f.size.pt for f in sources if f.size is not None), self.defaults["size"])
        font = None
        for f in sources:
            font = f.name or self._font_from_rfonts(f.element.rPr.rFonts if f.element.rPr is not None else None)
            if font:
                break
        out["font"] = font or self.defaults["font"]
        out["color"] = next((str(f.color.rgb) for f in sources if f.color is not None and f.color.type == 1
                             and f.color.rgb is not None), None)
        out["underline"] = bool(out["underline"])
        return out

    def para_fmt(self, para) -> dict | None:
        """Formatting of a paragraph's first visible run."""
        for run in para.runs:
            if run.text.strip():
                return self.fmt(run, para)
        return None

    def alignment(self, para) -> str:
        al = para.alignment
        if al is None:
            al = next((s.paragraph_format.alignment for s in self._chain(para.style)
                       if s.paragraph_format.alignment is not None), None)
        return "center" if al == 1 else "left"

    # ------------------------------------------------------------ structure
    def blocks(self):
        body = self.d.element.body
        for el in body.iterchildren():
            if el.tag == qn("w:p"):
                yield Paragraph(el, self.d)
            elif el.tag == qn("w:tbl"):
                yield Table(el, self.d)

    def is_bullet(self, para) -> bool:
        if para._p.pPr is not None and para._p.pPr.numPr is not None:
            return True
        return any(s.element.pPr is not None and s.element.pPr.numPr is not None for s in self._chain(para.style))

    def bullet_char(self, para) -> str | None:
        num_pr = para._p.pPr.numPr if para._p.pPr is not None else None
        if num_pr is None or num_pr.numId is None:
            return None
        numbering = self.d.part.numbering_part.element
        num = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_pr.numId.val)), None)
        if num is None:
            return None
        aid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == aid)
        lvl = abstract.find(qn("w:lvl"))
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is not None and fmt.get(qn("w:val")) != "bullet":
            return None
        text = lvl.find(qn("w:lvlText")).get(qn("w:val"))
        return BULLET_MAP.get(text, text if text and not ("" <= text[0] <= "") else "•")

    def not_carried_over(self) -> list[str]:
        xml = self.d.element.xml
        found = []
        if "<w:drawing" in xml or "<w:pict" in xml:
            found.append("images or shapes")
        if "txbxContent" in xml:
            found.append("text boxes")
        cols = [c for s in self.d.sections for c in s._sectPr.findall(qn("w:cols"))]
        if any(int(c.get(qn("w:num"), "1")) > 1 for c in cols):
            found.append("multi-column page layout")
        if len(self.d.sections) > 1:
            found.append("more than one page section")
        for s in self.d.sections:
            if any(p.text.strip() for p in s.header.paragraphs) or any(p.text.strip() for p in s.footer.paragraphs):
                found.append("header or footer content")
                break
        return found


def section_of(text: str) -> str | None:
    t = text.strip().lower()
    if not t or len(t.split()) > 6 or t.endswith(":") or YEAR_RANGE.search(t):
        return None
    for name, pattern in SECTION_WORDS:
        if re.search(pattern, t):
            return name
    return None


def style_line(f: dict, body: dict | None, heading: bool = False, rule: bool = False) -> str:
    """A Text value: font and size only where they differ from Body, then flags and colour."""
    parts = []
    if body is None or f["font"] != body["font"]:
        parts.append(f["font"])
    if body is None or abs(f["size"] - body["size"]) >= 0.25:
        parts.append(f"{f['size']:g}")
    parts += [name for name, key in (("bold", "bold"), ("italic", "italic"), ("caps", "all_caps"),
                                     ("underline", "underline")) if f[key]]
    if f["color"] and f["color"] not in ("000000", "auto") and (body is None or f["color"] != body.get("color")):
        parts.append(f"color {f['color']}")
    if heading and rule:
        parts.append("rule")
    return ", ".join(parts)


def extract(path: str) -> tuple[str, dict]:
    t = Template(path)
    guesses = []

    def guess(key, value, how, confidence="high"):
        guesses.append({"key": key, "value": value, "how": how, "confidence": confidence})

    blocks = list(t.blocks())
    paras = [b for b in blocks if isinstance(b, Paragraph) and b.text.strip()]
    if not paras:
        raise ValueError("the template has no text to learn from")

    # header: name, contact, headline
    name = paras[0]
    contact = next((p for p in paras[1:5] if "@" in p.text or re.search(r"[·•|]", p.text)), None)
    headline = next((p for p in paras[1:5] if p is not contact and len(p.text.split()) <= 12
                     and section_of(p.text) is None), None)
    roles = {"Name": t.para_fmt(name)}
    guess("Name", name.text.strip()[:40], "first line of the document")
    if contact is not None:
        roles["Contact"] = t.para_fmt(contact)
        guess("Contact", contact.text.strip()[:40], "line with an email address or separators near the top")
    if headline is not None:
        roles["Headline"] = t.para_fmt(headline)
        guess("Headline", headline.text.strip()[:40], "short line under the name", "medium")
    separator = None
    if contact is not None:
        m = re.search(r"\s*[·•|]\s*", contact.text)
        separator = f" {m.group(0).strip()} " if m else None

    # sections, in document order
    order, titles, heading_fmts, heading_rules, heading_paras = [], {}, [], [], []
    current, rows_after = None, collections.defaultdict(list)
    competencies = ("inline", " · ")
    for b in blocks:
        if isinstance(b, Table):
            if current == "competencies":
                cols = len(b.columns)
                competencies = ("grid", max(1, min(cols, 4)))
                guess("Competencies", f"grid {competencies[1]}", f"a {cols}-column table under the competencies heading")
            continue
        text = b.text.strip()
        sec = section_of(text) if b not in (name, contact, headline) else None
        if sec and (sec in titles or ":" in text or t.is_bullet(b)):
            sec = None  # the first heading for a section wins; later look-alikes are content
        if sec:
            both = sec == "education" and re.search(r"certif", text, re.I)
            for s in (["education", "certifications"] if both else [sec]):
                if s not in order:
                    order.append(s)
            titles[sec] = text
            if both:
                titles["certifications"] = ""
            current = sec
            heading_fmts.append(t.para_fmt(b))
            heading_paras.append(b)
            ppr = b._p.pPr
            heading_rules.append(ppr is not None and ppr.find(qn("w:pBdr")) is not None)
            continue
        if text:
            rows_after[current].append(b)

    # a long paragraph before the first section heading is the summary, shown without a heading
    first_heading_seen = False
    for b in blocks:
        if isinstance(b, Paragraph) and section_of(b.text.strip()):
            first_heading_seen = True
        if (not first_heading_seen and isinstance(b, Paragraph) and b not in (name, contact, headline)
                and len(b.text.split()) >= 12 and "summary" not in order):
            order.insert(0, "summary")
            titles["summary"] = ""
            rows_after["summary"].append(b)
            guess("Section order", "summary first, no heading", "a paragraph before the first section heading")
            break
    if heading_fmts:
        roles["Section headings"] = collections.Counter(
            json.dumps(f, sort_keys=True) for f in heading_fmts if f).most_common(1)[0][0]
        roles["Section headings"] = json.loads(roles["Section headings"])
        guess("Section headings", ", ".join(titles[s] or "(none)" for s in order),
              f"{len(heading_fmts)} lines naming a resume section")
    rule = sum(heading_rules) > len(heading_rules) / 2
    aligns = [t.alignment(p) for p in heading_paras]
    heading_align = collections.Counter(aligns).most_common(1)[0][0] if aligns else "left"
    position = {id(p._p): i for i, p in enumerate(paras)}
    present = [(position[id(p._p)], key) for key, p in (("name", name), ("headline", headline), ("contact", contact)) if p is not None]
    header_order = [k for _, k in sorted(present)]
    header_order += [k for k in ("name", "headline", "contact") if k not in header_order]

    # experience: company and role lines, descriptions, dates, bullets, label
    exp = rows_after.get("experience", [])
    dated = [p for p in exp if YEAR_RANGE.search(p.text) and not t.is_bullet(p)]
    company_lines, role_lines, descriptions, label = [], [], [], ""
    for i, p in enumerate(exp):
        if t.is_bullet(p):
            continue
        txt = p.text.strip()
        if re.fullmatch(r"[A-Z][\w &/-]{2,40}:", txt):
            label = label or txt
            continue
        f = t.para_fmt(p)
        nxt = exp[i + 1] if i + 1 < len(exp) else None
        nxt_f = t.para_fmt(nxt) if nxt is not None else None
        if p in dated and nxt_f and nxt_f["italic"] and not t.is_bullet(nxt):
            company_lines.append(p)
            descriptions.append(nxt)
        elif f and f["bold"] and len(txt.split()) <= 16:
            role_lines.append(p)
    if dated and not company_lines:
        company_lines = dated[:1]
    body_samples = [p for p in exp + rows_after.get("summary", []) + rows_after.get("highlights", [])
                    if p not in company_lines + role_lines + descriptions and len(p.text.split()) >= 6]
    body_fmts = [t.para_fmt(p) for p in body_samples]
    body_fmts = [f for f in body_fmts if f]
    body = (json.loads(collections.Counter(json.dumps({"font": f["font"], "size": f["size"], "color": f["color"]},
                                                      sort_keys=True) for f in body_fmts).most_common(1)[0][0])
            if body_fmts else {"font": t.defaults["font"], "size": t.defaults["size"], "color": None})
    guess("Body", f"{body['font']}, {body['size']:g}", f"most common format of {len(body_fmts)} body paragraphs",
          "high" if len(body_fmts) >= 3 else "low")
    if company_lines:
        roles["Company"] = t.para_fmt(company_lines[0])
        guess("Company", company_lines[0].text.split("\t")[0].strip()[:40],
              "dated line followed by an italic description", "medium")
    if descriptions:
        roles["Company description"] = t.para_fmt(descriptions[0])
    # location: the first run on a company or role line after the name, that isn't a separator or the dates
    for line in company_lines[:1] + role_lines[:1]:
        runs = [r for r in line.runs if r.text.strip(" ,|\t")]
        loc = next((r for r in runs[1:] if not re.search(r"(19|20)\d{2}|present", r.text, re.I)), None)
        if loc is not None:
            roles["Location"] = t.fmt(loc, line)
            break
    if role_lines:
        roles["Role"] = t.para_fmt(role_lines[0])
        guess("Role", role_lines[0].text.split("\t")[0].strip()[:40], "bold line in the experience section", "medium")
    dates_right = any("\t" in p.text for p in dated)
    if dated:
        guess("Dates", "right" if dates_right else "inline",
              "tab before the dates" if dates_right else "no tab before the dates")
        date_run = next((r for r in dated[0].runs if YEAR_RANGE.search(r.text) or re.search(r"(19|20)\d{2}", r.text)), None)
        if date_run is not None:
            roles["Dates"] = t.fmt(date_run, dated[0])
    bullets = [b for b in blocks if isinstance(b, Paragraph) and t.is_bullet(b)]
    bullet = next((c for c in (t.bullet_char(b) for b in bullets) if c), "•")
    if bullets:
        guess("Bullet", bullet, "the template's list definition")
    if label:
        guess("Accomplishments label", label, "a bold line ending in ':' before bullets")

    # page
    sec = t.d.sections[0]
    w, h = sec.page_width.inches, sec.page_height.inches
    size = "A4" if abs(w - 8.27) < 0.1 and abs(h - 11.69) < 0.1 else "Letter"
    margins = {k: round(getattr(sec, f"{k}_margin").inches, 2) for k in ("top", "bottom", "left", "right")}
    margins = {k: min(max(v, 0.2), 2.0) for k, v in margins.items()}

    # fonts the renderer can't count on
    nco = t.not_carried_over()
    fonts = {f["font"] for f in list(roles.values()) + [body] if isinstance(f, dict) and f.get("font")}
    for font in sorted(fonts):
        if font.lower() not in COMMON_FONTS:
            nco.append(f"font '{font}' may not be installed where PDFs are made; a similar font is used there")

    # write the style file
    text_lines = [f"- Body: {body['font']}, {body['size']:g}" + (f", color {body['color']}" if body.get("color") and body["color"] != "000000" else "")]
    for role in ("Name", "Headline", "Contact", "Section headings", "Company", "Company description", "Role",
                 "Location", "Dates"):
        f = roles.get(role)
        if f:
            value = style_line(f, body, heading=role == "Section headings", rule=rule)
            if value:
                text_lines.append(f"- {role}: {value}")
    titles_value = "; ".join(f"{s} = {titles.get(s, '')}" for s in order)
    default_order = ["summary", "competencies", "highlights", "experience", "education", "certifications", "additional"]
    order = order or default_order
    out = [
        "# Resume style",
        "Style format: 1",
        f"Source: {Path(path).name}, converted {datetime.date.today().isoformat()}",
        "",
        "## Page",
        f"- Size: {size}",
        "- Margins (in): " + ", ".join(f"{k} {v:g}" for k, v in margins.items()),
        "",
        "## Text",
        *text_lines,
        "",
        "## Layout",
        f"- Header alignment: {t.alignment(name)}",
        f"- Header order: {', '.join(header_order)}",
        f"- Section heading alignment: {heading_align}",
        f'- Contact separator: "{separator or " · "}"',
        f"- Section order: {', '.join(order)}",
        f"- Section titles: {titles_value}",
        f"- Competencies: grid {competencies[1]}" if competencies[0] == "grid" else f'- Competencies: inline "{competencies[1]}"',
        f"- Bullet: {bullet}",
        f"- Dates: {'right' if dates_right else 'inline'}",
        f"- Accomplishments label: {label}",
        "",
        "## Not carried over",
        *([f"- {x}" for x in nco] or ["- Nothing detected."]),
        "",
    ]
    return "\n".join(out), {"guesses": guesses, "not_carried_over": nco}


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("template")
    a.add_argument("--out")
    a = a.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        style, report = extract(a.template)
    except Exception as e:  # noqa: BLE001  any unreadable input is reported, not raised
        print(json.dumps({"error": f"couldn't read {a.template} as a Word resume: {e}"}))
        sys.exit(2)
    if a.out:
        Path(a.out).write_text(style, encoding="utf-8")
    else:
        print(style)
    print(json.dumps({"style": a.out, **report}, ensure_ascii=False), file=sys.stderr if not a.out else sys.stdout)


if __name__ == "__main__":
    main()
