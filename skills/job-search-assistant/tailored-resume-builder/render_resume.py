#!/usr/bin/env python3
"""Render a resume .docx (and PDF) from content.json and a resume style file.

Usage: render_resume.py --style resume-style.md --content content.json --out out.docx [--pdf] [--max-pages N]
                        [--html out.html]

The style file (format: resume-style-format.md; example: default-style.md) is the user's look,
as text: page, fonts, colours, section order and layout. The document is built from scratch, so
no Word template is needed. Prints JSON: {"out": ..., "pages": N or null, "max_pages": N,
"last_page_fill": 0-1 or null, "warnings": [...]}. The page limit is 3 unless --max-pages lowers it,
and a 3rd page must be at least half full. Exit 2, with {"error": ...}, on an invalid style file,
content.json or --max-pages.

--html also writes the resume as HTML for Google Drive to convert into a Google Doc (see HtmlResume).
"""

from __future__ import annotations

import argparse
import html
import json
import math
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field, replace

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

STYLE_FORMAT = "1"
SECTIONS = ("summary", "competencies", "highlights", "experience", "education", "certifications", "additional")
DEFAULT_TITLES = {
    "summary": "Summary", "competencies": "Core Competencies", "highlights": "Career Highlights",
    "experience": "Professional Experience", "education": "Education",
    "certifications": "Certifications", "additional": "Additional Information",
}
TEXT_ROLES = ("Body", "Name", "Headline", "Contact", "Section headings", "Company",
              "Company description", "Role", "Location", "Dates")
PAGE_SIZES = {"letter": (8.5, 11.0), "a4": (8.27, 11.69)}
LIMITS = dict(competencies=12, highlights=5)
MAX_PAGES = 3  # the hard limit; the user can ask for fewer
MIN_LAST_PAGE_FILL = 0.5  # a 3rd page must be at least this full, or the resume goes back to 2 pages


class StyleError(ValueError):
    pass


def fail(msg: str):
    print(json.dumps({"error": msg}))
    sys.exit(2)


# ---------------------------------------------------------------- style file


@dataclass
class TextStyle:
    font: str | None = None
    size: float | None = None
    bold: bool = False
    italic: bool = False
    caps: bool = False
    underline: bool = False
    color: str | None = None
    rule: bool = False


@dataclass
class Style:
    size: str = "letter"
    margins: dict = field(default_factory=lambda: dict(top=0.5, bottom=0.5, left=0.6, right=0.6))
    text: dict = field(default_factory=dict)
    align: str = "left"
    header_order: list = field(default_factory=lambda: ["name", "headline", "contact"])
    heading_align: str = "left"
    contact_sep: str = " · "
    order: list = field(default_factory=lambda: list(SECTIONS))
    titles: dict = field(default_factory=lambda: dict(DEFAULT_TITLES))
    competencies: tuple = ("grid", 3)
    bullet: str = "•"
    dates: str = "right"
    acc_label: str = ""
    space_heading: float = 3.0
    space_roles: float = 6.0

    def t(self, role: str) -> TextStyle:
        """The text style for a role, with font, size and colour inherited from Body."""
        body, own = self.text["Body"], self.text.get(role, TextStyle())
        return replace(own, font=own.font or body.font, size=own.size or body.size, color=own.color or body.color)


def unquote(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
        return value[1:-1]
    return value


def parse_text_style(value: str, where: str, heading: bool) -> TextStyle:
    ts = TextStyle()
    for token in (t.strip() for t in value.split(",")):
        low = token.lower()
        if not token:
            continue
        if re.fullmatch(r"\d+(\.\d+)?", token):
            ts.size = float(token)
            if not 4 <= ts.size <= 72:
                raise StyleError(f"{where}: size {token} is out of range (4-72)")
        elif low in ("bold", "italic", "caps", "underline"):
            setattr(ts, low, True)
        elif low == "rule":
            if not heading:
                raise StyleError(f"{where}: 'rule' only applies to Section headings")
            ts.rule = True
        elif low.startswith("color"):
            hex_ = token[5:].strip().lstrip("#")
            if not re.fullmatch(r"[0-9A-Fa-f]{6}", hex_):
                raise StyleError(f"{where}: color must be 6 hex digits, e.g. color 1F3864")
            ts.color = hex_.upper()
        elif ts.font is None and re.fullmatch(r"[A-Za-z][\w .'-]*", token):
            ts.font = token
        else:
            raise StyleError(f"{where}: don't understand '{token}'")
    return ts


def parse_style(text: str) -> Style:
    st = Style()
    section, version, seen_text = None, None, {}
    for n, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if line.startswith("## "):
            section = line[3:].strip().lower()
            if section not in ("page", "text", "layout", "not carried over"):
                raise StyleError(f"line {n}: unknown section '{line[3:].strip()}'")
            continue
        if section is None:
            m = re.match(r"Style format:\s*(\S+)", line)
            if m:
                version = m.group(1)
            continue
        if section == "not carried over" or not line.startswith("- "):
            continue
        key, sep, value = line[2:].partition(":")
        if not sep:
            raise StyleError(f"line {n}: expected '- Key: value'")
        key, value, where = key.strip(), value.strip(), f"line {n} ({key.strip()})"
        if section == "page":
            if key == "Size":
                if value.lower() not in PAGE_SIZES:
                    raise StyleError(f"{where}: size must be Letter or A4")
                st.size = value.lower()
            elif key == "Margins (in)":
                st.margins = parse_margins(value, where)
            else:
                raise StyleError(f"{where}: unknown Page setting")
        elif section == "text":
            if key not in TEXT_ROLES:
                raise StyleError(f"{where}: unknown Text setting (expected one of: {', '.join(TEXT_ROLES)})")
            seen_text[key] = parse_text_style(value, where, heading=key == "Section headings")
        elif section == "layout":
            parse_layout(st, key, value, where)
    if version != STYLE_FORMAT:
        raise StyleError(f"'Style format: {STYLE_FORMAT}' is required (found {version!r})")
    body = seen_text.get("Body")
    if body is None or body.font is None or body.size is None:
        raise StyleError("Text → Body must give a font and a size, e.g. 'Body: Calibri, 10.5'")
    st.text = seen_text
    return st


def parse_margins(value: str, where: str) -> dict:
    if re.fullmatch(r"\d+(\.\d+)?", value):
        return dict.fromkeys(("top", "bottom", "left", "right"), float(value))
    margins = {}
    for part in value.split(","):
        m = re.fullmatch(r"\s*(top|bottom|left|right)\s+(\d+(?:\.\d+)?)\s*", part)
        if not m:
            raise StyleError(f"{where}: expected 'top 0.5, bottom 0.5, left 0.6, right 0.6'")
        margins[m.group(1)] = float(m.group(2))
    if set(margins) != {"top", "bottom", "left", "right"}:
        raise StyleError(f"{where}: give all four margins")
    if any(not 0.2 <= v <= 2 for v in margins.values()):
        raise StyleError(f"{where}: margins must be between 0.2 and 2 inches")
    return margins


def parse_layout(st: Style, key: str, value: str, where: str) -> None:
    if key in ("Header alignment", "Section heading alignment"):
        if value.lower() not in ("left", "center"):
            raise StyleError(f"{where}: left or center")
        setattr(st, "align" if key.startswith("Header") else "heading_align", value.lower())
    elif key == "Header order":
        order = [s.strip().lower() for s in value.split(",") if s.strip()]
        if sorted(order) != ["contact", "headline", "name"]:
            raise StyleError(f"{where}: name, headline and contact, in any order")
        st.header_order = order
    elif key == "Contact separator":
        st.contact_sep = unquote(value)
    elif key == "Section order":
        order = [s.strip().lower() for s in value.split(",") if s.strip()]
        bad = [s for s in order if s not in SECTIONS]
        if bad or len(set(order)) != len(order):
            raise StyleError(f"{where}: sections must be distinct, from: {', '.join(SECTIONS)}")
        st.order = order
    elif key == "Section titles":
        for pair in filter(None, (p.strip() for p in value.split(";"))):
            name, eq, title = pair.partition("=")
            name = name.strip().lower()
            if not eq or name not in SECTIONS:
                raise StyleError(f"{where}: expected 'section = Title' pairs; '{pair}' isn't one")
            st.titles[name] = unquote(title)
    elif key == "Competencies":
        m = re.fullmatch(r"grid\s+([1-4])", value.strip(), re.I)
        if m:
            st.competencies = ("grid", int(m.group(1)))
        elif value.lower().startswith("inline"):
            st.competencies = ("inline", unquote(value[6:]) or " · ")
        else:
            raise StyleError(f"{where}: 'grid N' (1-4) or 'inline \" · \"'")
    elif key == "Bullet":
        if not value or len(unquote(value)) > 2:
            raise StyleError(f"{where}: one bullet character, e.g. •")
        st.bullet = unquote(value)
    elif key == "Dates":
        if value.lower() not in ("right", "inline"):
            raise StyleError(f"{where}: right or inline")
        st.dates = value.lower()
    elif key == "Accomplishments label":
        st.acc_label = unquote(value)
    elif key in ("Space after section heading (pt)", "Space between roles (pt)"):
        if not re.fullmatch(r"\d+(\.\d+)?", value):
            raise StyleError(f"{where}: a number of points")
        setattr(st, "space_heading" if "heading" in key else "space_roles", float(value))
    else:
        raise StyleError(f"{where}: unknown Layout setting")


# ---------------------------------------------------------------- content


def content_problems(c: dict) -> list[str]:
    """Everything wrong with content.json, as messages; empty when it's valid."""
    missing = [k for k in ("name", "contact", "headline", "summary", "competencies", "highlights",
                           "experience", "education") if k not in c]
    if missing:
        return [f"content.json missing '{k}'" for k in missing]
    problems = []
    if len(c["competencies"]) > LIMITS["competencies"]:
        problems.append("too many competencies (max 12)")
    if len(c["highlights"]) > LIMITS["highlights"]:
        problems.append("too many career highlights (max 5)")
    for j in c["experience"]:
        absent = [k for k in ("company", "dates", "roles") if k not in j]
        if absent:
            problems += [f"experience entry missing '{k}'" for k in absent]
            continue
        if j.get("description"):
            problems.append(f"company descriptions aren't shown on tailored resumes; "
                            f"remove 'description' from {j['company']}")
        for r in j["roles"]:
            absent = [k for k in ("title", "dates") if k not in r]
            if absent:
                problems += [f"role missing '{k}'" for k in absent]
            elif not r.get("bullets") and not (r.get("summary") or "").strip():
                problems.append(f"{r['title']} at {j['company']} needs at least one line (a summary or a bullet)")
    return problems


def validate(c: dict) -> None:
    problems = content_problems(c)
    if problems:
        fail(problems[0])


def normalize(text: str) -> str:
    """Dates and places compared loosely: any dash, any spacing, any case."""
    return re.sub(r"\s+", " ", re.sub(r"\s*[-\u2013\u2014]\s*", " - ", text or "")).strip().lower()


def same(a: str, b: str) -> bool:
    return bool(a) and normalize(a) == normalize(b)


# ---------------------------------------------------------------- rendering


class Renderer:
    def __init__(self, st: Style):
        self.st = st
        self.d = docx.Document()
        body = self.d.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        sec = self.d.sections[0]
        width, height = PAGE_SIZES[st.size]
        sec.page_width, sec.page_height = Inches(width), Inches(height)
        for side, inches in st.margins.items():
            setattr(sec, f"{side}_margin", Inches(inches))
        self.text_width = Inches(width - st.margins["left"] - st.margins["right"])
        normal = self.d.styles["Normal"]
        b = st.t("Body")
        normal.font.name, normal.font.size = b.font, Pt(b.size)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), b.font)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        self.num_id = self._bullet_numbering()

    # -- low-level helpers
    def run(self, p, text: str, role: str, **override):
        ts = replace(self.st.t(role), **override)
        r = p.add_run(text)
        f = r.font
        f.name, f.size = ts.font, Pt(ts.size)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), ts.font)
        f.bold, f.italic, f.all_caps, f.underline = ts.bold or None, ts.italic or None, ts.caps or None, ts.underline or None
        if ts.color:
            f.color.rgb = RGBColor.from_string(ts.color)
        return r

    def para(self, after: float = 0.0, align: str | None = None):
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def _bullet_numbering(self) -> int:
        """Add a one-level bullet list using the style's bullet character; return its numId."""
        numbering = self.d.part.numbering_part.element
        abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        aid, nid = max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(aid))
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", self.st.bullet), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            lvl.append(el)
        ppr, ind = OxmlElement("w:pPr"), OxmlElement("w:ind")
        ind.set(qn("w:left"), "300")
        ind.set(qn("w:hanging"), "200")
        ppr.append(ind)
        lvl.append(ppr)
        rpr, fonts = OxmlElement("w:rPr"), OxmlElement("w:rFonts")
        body_font = self.st.t("Body").font
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), body_font)
        rpr.append(fonts)
        lvl.append(rpr)
        a.append(lvl)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(aid))
        num.append(ref)
        # abstractNum elements must come before num elements
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            numbering.append(a)
        numbering.append(num)
        return nid

    def bullet(self, text: str, container=None):
        p = container.add_paragraph() if container is not None else self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = self.num_id
        self.run(p, text, "Body")
        return p

    def heading(self, section: str):
        title = self.st.titles.get(section, DEFAULT_TITLES[section])
        if not title:
            return
        p = self.para(after=self.st.space_heading, align=self.st.heading_align)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        self.run(p, title, "Section headings")
        if self.st.t("Section headings").rule:
            color = self.st.t("Section headings").color or "auto"
            bdr, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
                bottom.set(qn(k), v)
            bdr.append(bottom)
            p._p.get_or_add_pPr().append(bdr)

    def dated_line(self, left_role: str, left: str, location: str, dates: str, after: float = 0.0):
        p = self.para(after=after)
        p.paragraph_format.keep_with_next = True
        self.run(p, left, left_role)
        if location:
            self.run(p, ", ", left_role, bold=False, italic=False, caps=False)
            self.run(p, location, "Location")
        if dates:
            if self.st.dates == "right":
                p.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
                p.add_run("\t")
            else:
                self.run(p, " | ", "Dates")
            self.run(p, dates, "Dates")
        return p

    # -- sections
    def header(self, c: dict):
        a = self.st.align
        contact = self.st.contact_sep.join(x for x in c.get("contact", []) if x)
        lines = {"name": (c["name"], "Name", 0), "headline": (c.get("headline"), "Headline", 1),
                 "contact": (contact, "Contact", 1)}
        for i, part in enumerate(self.st.header_order):
            text, role, after = lines[part]
            if text:
                self.run(self.para(after=2 if i == 2 else after, align=a), text, role)

    def summary(self, c):
        if c.get("summary"):
            self.heading("summary")
            self.run(self.para(after=2), c["summary"], "Body")

    def competencies(self, c):
        items = c.get("competencies") or []
        if not items:
            return
        self.heading("competencies")
        kind, arg = self.st.competencies
        if kind == "inline":
            self.run(self.para(after=2), arg.join(items), "Body")
            return
        cols = min(arg, len(items))
        rows = math.ceil(len(items) / cols)
        table = self.d.add_table(rows=1, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, cell in enumerate(table.rows[0].cells):
            chunk = items[i * rows:(i + 1) * rows]
            cell._tc.remove(cell.paragraphs[0]._p)
            for item in chunk:
                self.bullet(item, container=cell)
            if not chunk:
                cell.add_paragraph()

    def highlights(self, c):
        if c.get("highlights"):
            self.heading("highlights")
            for h in c["highlights"]:
                self.bullet(h)

    def experience(self, c):
        if not c.get("experience"):
            return
        self.heading("experience")
        gap = Pt(self.st.space_roles)
        for j, job in enumerate(c["experience"]):
            line = self.dated_line("Company", job["company"], job.get("location", ""), job["dates"])
            if j:
                line.paragraph_format.space_before = gap
            for i, role in enumerate(job["roles"]):
                # a role's place and dates are shown only where they differ from the employer's
                location = "" if same(role.get("location", ""), job.get("location", "")) else role.get("location", "")
                dates = "" if same(role["dates"], job["dates"]) else role["dates"]
                line = self.dated_line("Role", role["title"], location, dates, after=1)
                if i:  # the space goes before each later role, so no empty spacer paragraphs
                    line.paragraph_format.space_before = gap
                if role.get("summary"):
                    self.run(self.para(after=1), role["summary"], "Body")
                bullets = role.get("bullets") or []
                if bullets and self.st.acc_label:
                    self.run(self.para(), self.st.acc_label, "Body", bold=True)
                for b in bullets:
                    self.bullet(b)

    def _pairs(self, section, items, first, second):
        if not items:
            return
        self.heading(section)
        for it in items:
            p = self.para(after=1)
            self.run(p, it[first], "Body", bold=True)
            if it.get(second):
                self.run(p, ", " + it[second], "Body")

    def education(self, c):
        self._pairs("education", c.get("education"), "degree", "school")

    def certifications(self, c):
        self._pairs("certifications", c.get("certifications"), "name", "detail")

    def additional(self, c):
        if c.get("additional"):
            self.heading("additional")
            for line in c["additional"]:
                self.run(self.para(after=1), line, "Body")

    def render(self, c: dict) -> None:
        self.header(c)
        for section in self.st.order:
            getattr(self, section)(c)

    def save(self, out: str) -> None:
        self.d.save(out)


# ---------------------------------------------------------------- Google Doc (HTML)


class HtmlResume:
    """The same resume as HTML that Google Drive converts into a Google Doc (create_file with
    contentMimeType text/html). Built for Drive's converter, tested through it:
    - page margins go on the body's padding, the way Google's own HTML export carries them;
    - dates sit in the right cell of a borderless two-column table, because Docs has no tab
      stops (Docs caps an imported table at its default width, so the dates column can end a
      little short of the right margin);
    - a company line and its first role share one table, since Docs puts an empty paragraph
      between two tables;
    - a heading rule is an <hr> followed by a 1pt spacer, because an <hr> right before a list
      becomes an empty bullet;
    - styles are classes, which keeps the text short enough to pass through the connector."""

    DATES_PT = 160

    def __init__(self, st: Style):
        self.st = st
        self.classes: dict[str, str] = {}
        self.text_pt = (PAGE_SIZES[st.size][0] - st.margins["left"] - st.margins["right"]) * 72

    def cls(self, css: str) -> str:
        self.classes.setdefault(css, f"c{len(self.classes)}")
        return f'class="{self.classes[css]}"'

    def font(self, role: str, **override) -> str:
        t = replace(self.st.t(role), **override)
        parts = [f"font-family:'{t.font}'", f"font-size:{t.size:g}pt"]
        parts += ["font-weight:bold"] * t.bold + ["font-style:italic"] * t.italic
        parts += ["text-transform:uppercase"] * t.caps + ["text-decoration:underline"] * t.underline
        if t.color:
            parts.append(f"color:#{t.color}")
        return self.cls(";".join(parts))

    def span(self, text: str, role: str, **override) -> str:
        return f"<span {self.font(role, **override)}>{html.escape(text)}</span>" if text else ""

    def p(self, inner: str, after: float = 0, before: float = 0, align: str | None = None) -> str:
        css = f"margin:{before:g}pt 0 {after:g}pt 0" + (f";text-align:{align}" if align == "center" else "")
        return f"<p {self.cls(css)}>{inner}</p>"

    def bullets(self, items) -> str:
        lis = "".join(f'<li {self.cls("margin:0 0 1pt 0")}>{self.span(i, "Body")}</li>' for i in items)
        return f'<ul {self.cls("margin:0;padding-left:18pt")}>{lis}</ul>'

    def heading(self, section: str) -> str:
        title = self.st.titles.get(section, DEFAULT_TITLES[section])
        if not title:
            return ""
        ts = self.st.t("Section headings")
        out = self.p(self.span(title, "Section headings"), before=8, after=0 if ts.rule else self.st.space_heading,
                     align=self.st.heading_align)
        if ts.rule:
            rule = f"border:none;border-top:1px solid #{ts.color or '000000'};margin:0"
            out += f"<hr {self.cls(rule)}>"
            out += f'<p {self.cls("margin:0;font-size:1pt;line-height:1pt")}>&nbsp;</p>'
        return out

    def dated_row(self, left: str, role: str, location: str, dates: str) -> str:
        inner = self.span(left, role)
        if location:
            inner += self.span(", ", role, bold=False, italic=False, caps=False) + self.span(location, "Location")
        left_w, right_w = self.text_pt - self.DATES_PT, self.DATES_PT
        if self.st.dates == "inline" and dates:
            inner += self.span(" | ", "Dates") + self.span(dates, "Dates")
            dates = ""
        return (f'<tr><td {self.cls(f"border:none;padding:0;width:{left_w:g}pt")}>{inner}</td>'
                f'<td {self.cls(f"border:none;padding:0;width:{right_w:g}pt;text-align:right")}>'
                f'{self.span(dates, "Dates")}</td></tr>')

    def table(self, rows: list, before: float) -> str:
        return f'<table {self.cls(f"width:{self.text_pt:g}pt;border-collapse:collapse;margin-top:{before:g}pt")}>' \
               f'{"".join(rows)}</table>'

    def section(self, name: str, c: dict) -> str:
        st, out = self.st, []
        if name == "summary" and c.get("summary"):
            out += [self.heading("summary"), self.p(self.span(c["summary"], "Body"), after=2)]
        elif name == "competencies" and c.get("competencies"):
            out.append(self.heading("competencies"))
            kind, arg = st.competencies
            items = c["competencies"]
            if kind == "inline":
                out.append(self.p(self.span(arg.join(items), "Body"), after=2))
            else:
                cols = min(arg, len(items))
                rows = math.ceil(len(items) / cols)
                cells = "".join(f'<td {self.cls("border:none;padding:0;vertical-align:top")}>'
                                f'{self.bullets(items[i * rows:(i + 1) * rows])}</td>' for i in range(cols))
                out.append(f'<table {self.cls("width:100%;border-collapse:collapse")}><tr>{cells}</tr></table>')
        elif name == "highlights" and c.get("highlights"):
            out += [self.heading("highlights"), self.bullets(c["highlights"])]
        elif name == "experience" and c.get("experience"):
            out.append(self.heading("experience"))
            for j, job in enumerate(c["experience"]):
                rows = [self.dated_row(job["company"], "Company", job.get("location", ""), job["dates"])]
                before = st.space_roles if j else 0
                for role in job["roles"]:
                    location = "" if same(role.get("location", ""), job.get("location", "")) else role.get("location", "")
                    dates = "" if same(role["dates"], job["dates"]) else role["dates"]
                    rows.append(self.dated_row(role["title"], "Role", location, dates))
                    out.append(self.table(rows, before))
                    rows, before = [], st.space_roles
                    if role.get("summary"):
                        out.append(self.p(self.span(role["summary"], "Body"), after=1))
                    bullets = role.get("bullets") or []
                    if bullets and st.acc_label:
                        out.append(self.p(self.span(st.acc_label, "Body", bold=True)))
                    if bullets:
                        out.append(self.bullets(bullets))
        elif name in ("education", "certifications") and c.get(name):
            out.append(self.heading(name))
            first, second = ("degree", "school") if name == "education" else ("name", "detail")
            for it in c[name]:
                rest = self.span(", " + it[second], "Body") if it.get(second) else ""
                out.append(self.p(self.span(it[first], "Body", bold=True) + rest, after=1))
        elif name == "additional" and c.get("additional"):
            out.append(self.heading("additional"))
            out += [self.p(self.span(line, "Body"), after=1) for line in c["additional"]]
        return "".join(out)

    def render(self, c: dict) -> str:
        st = self.st
        contact = st.contact_sep.join(x for x in c.get("contact", []) if x)
        lines = {"name": (c["name"], "Name"), "headline": (c.get("headline"), "Headline"),
                 "contact": (contact, "Contact")}
        body = []
        for part in st.header_order:
            text, role = lines[part]
            if text:
                body.append(self.p(self.span(text, role), after=2, align=st.align))
        body += [self.section(name, c) for name in st.order]
        m = {k: v * 72 for k, v in st.margins.items()}
        b = st.t("Body")
        page = (f"body{{font-family:'{b.font}';font-size:{b.size:g}pt;max-width:{self.text_pt:g}pt;"
                f"padding:{m['top']:g}pt {m['right']:g}pt {m['bottom']:g}pt {m['left']:g}pt}}")
        rules = "".join(f".{name}{{{css}}}" for css, name in self.classes.items())
        return (f'<html><head><meta charset="utf-8"><style>{page}{rules}</style></head>'
                f'<body>{"".join(body)}</body></html>')


def bbox_fill(bbox_html: str, margins: dict):
    """How far down the text area a page's last line reaches (0-1), from `pdftotext -bbox` output
    for that one page. Coordinates are points from the top of the page."""
    page = re.search(r'<page width="([\d.]+)" height="([\d.]+)"', bbox_html)
    bottoms = [float(y) for y in re.findall(r'<word [^>]*yMax="([\d.]+)"', bbox_html)]
    if not page or not bottoms:
        return None
    top, bottom = margins["top"] * 72, margins["bottom"] * 72
    usable = float(page.group(2)) - top - bottom
    return round(min(1.0, max(0.0, (max(bottoms) - top) / usable)), 2)


def last_page_fill(pdf: str, page: int, margins: dict):
    try:
        out = subprocess.run(["pdftotext", "-f", str(page), "-l", str(page), "-bbox", pdf, "-"],
                             check=True, capture_output=True, text=True, timeout=60).stdout
        return bbox_fill(out, margins)
    except Exception:
        return None


def count_pages(path: str, margins: dict):
    """(pages, PDF bytes, last-page fill), or Nones when LibreOffice or poppler aren't available."""
    try:
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, path],
                           check=True, capture_output=True, timeout=120)
            pdf = os.path.join(td, os.path.splitext(os.path.basename(path))[0] + ".pdf")
            out = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
            n = int([line for line in out.splitlines() if line.startswith("Pages:")][0].split()[1])
            fill = last_page_fill(pdf, n, margins)
            with open(pdf, "rb") as fh:
                return n, fh.read(), fill
    except Exception:
        return None, None, None


def page_warnings(pages, max_pages: int, fill) -> list[str]:
    if pages is None:
        return ["could not render to PDF; page count unverified"]
    if pages > max_pages:
        return [f"resume is {pages} pages; the limit is {max_pages}"]
    if pages == 3 and fill is None:
        return ["page 3 fill couldn't be measured; trim to 2 pages"]
    if pages == 3 and fill < MIN_LAST_PAGE_FILL:
        return [f"page 3 is only {round(fill * 100)}% full (needs {round(MIN_LAST_PAGE_FILL * 100)}%+); "
                "trim to 2 pages"]
    return []


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("--style", required=True)
    a.add_argument("--content", required=True)
    a.add_argument("--out", required=True)
    a.add_argument("--pdf", action="store_true")
    a.add_argument("--max-pages", type=int, default=MAX_PAGES)
    a.add_argument("--html", help="also write Google Doc HTML here")
    a = a.parse_args()
    if not 1 <= a.max_pages <= MAX_PAGES:
        fail(f"max pages must be 1, 2 or {MAX_PAGES}")
    try:
        with open(a.style, encoding="utf-8") as fh:
            st = parse_style(fh.read())
    except (OSError, StyleError) as e:
        fail(f"style file: {e}")
    with open(a.content, encoding="utf-8") as fh:
        c = json.load(fh)
    validate(c)
    r = Renderer(st)
    r.render(c)
    r.save(a.out)
    if a.html:
        with open(a.html, "w", encoding="utf-8") as fh:
            fh.write(HtmlResume(st).render(c))
    pages, pdf, fill = count_pages(a.out, st.margins)
    if a.pdf and pdf:
        with open(os.path.splitext(a.out)[0] + ".pdf", "wb") as fh:
            fh.write(pdf)
    result = {"out": a.out, "pages": pages, "max_pages": a.max_pages, "last_page_fill": fill,
              "warnings": page_warnings(pages, a.max_pages, fill)}
    if a.html:
        result["html"] = a.html
    print(json.dumps(result))


if __name__ == "__main__":
    main()
