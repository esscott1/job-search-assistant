#!/usr/bin/env python3
"""Make fixtures/Synthetic-Resume-Template.docx, a stand-in for the resume template.

build_resume.py clones prototype paragraphs from the template by their position in the
document body (its PROTO table). The real template isn't in the repo, so tests and evals
use this generated file instead: it has the same structure at the same positions, with
placeholder text only. Needs python-docx.

    python make_synthetic_template.py
"""

from __future__ import annotations

import datetime
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "tailored-resume-builder"))
from build_resume import PROTO  # noqa: E402

OUT = HERE / "fixtures" / "Synthetic-Resume-Template.docx"
BODY_LENGTH = max(PROTO.values()) + 1

# Prototype role -> (style, placeholder text, bold, italic, list item)
ROLES = {
    "title": ("Title", "SAMPLE NAME", True, False, False),
    "contact": ("Normal", "555-0000 · Sample City · sample@example.com", False, False, False),
    "headline": ("Subtitle", "SAMPLE HEADLINE", True, False, False),
    "summary": ("Normal", "Sample summary text.", False, False, False),
    "section": ("Heading 1", "SAMPLE SECTION", True, False, False),
    "highlight": ("List Bullet", "Sample highlight.", False, False, True),
    "company": ("Heading 2", "Sample Company", True, False, False),
    "company_desc": ("Normal", "Sample company description.", False, True, False),
    "role": ("Heading 3", "Sample Role", True, False, False),
    "role_summary": ("Normal", "Sample role summary.", False, False, False),
    "key_acc": ("Normal", "Key Accomplishments:", True, False, False),
    "bullet": ("List Bullet", "Sample accomplishment.", False, False, True),
    "spacer": ("Normal", "", False, False, False),
    "edu_title": ("Heading 1", "Education and Certifications", True, False, False),
    "degree": ("Normal", "Sample Degree", True, False, False),
    "school": ("Normal", "Sample School", False, False, False),
    "addl_title": ("Heading 1", "Additional Information", True, False, False),
    "addl": ("Normal", "Sample additional line.", False, False, False),
}


def numbered(paragraph) -> None:
    """Give the paragraph its own numPr, as list items in the real template have."""
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = 1


def main() -> int:
    d = docx.Document()
    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5))
    d.styles["Normal"].font.size = Pt(10)

    by_index = {i: role for role, i in PROTO.items()}
    for i in range(BODY_LENGTH):
        role = by_index.get(i)
        if role == "table":
            table = d.add_table(rows=1, cols=3)
            for cell in table.rows[0].cells:
                p = cell.paragraphs[0]
                p.style = d.styles["List Bullet"]
                p.add_run("Sample competency")
                numbered(p)
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            height = OxmlElement("w:trHeight")
            height.set(qn("w:val"), "400")
            tr_pr.append(height)
            continue
        style, text, bold, italic, listed = ROLES.get(role, ("Normal", "Sample filler text.", False, False, False))
        p = d.add_paragraph(style=d.styles[style])
        if text:
            run = p.add_run(text)
            run.bold, run.italic = bold or None, italic or None
            run.font.size = Pt(10)  # gives every text run an rPr to clone, as in the real template
        if listed:
            numbered(p)

    fixed = datetime.datetime(2026, 1, 1)
    d.core_properties.author = "job-search-assistant evals"
    d.core_properties.title = "Synthetic resume template"
    d.core_properties.created = d.core_properties.modified = fixed
    d.save(OUT)
    print(f"wrote {OUT.relative_to(HERE)} ({len(body) - 1} body elements)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
